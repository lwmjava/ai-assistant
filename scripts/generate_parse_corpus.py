"""生成 RAG-033「多格式解析与 OCR 质量基线」真实样本集。

样本全部使用本机**已安装依赖**真实生成（PyMuPDF / python-docx / openpyxl /
python-pptx），不手写二进制，不使用 Mock。生成结果落在
``evals/corpus/parsing/``，并在同目录写出 ``manifest.json`` 作为期望标注。

用法::

    python scripts/generate_parse_corpus.py

中文 PDF 需要系统 CJK 字体（脚本会按候选路径探测）。若探测不到则回退到
PyMuPDF 内置 ``china-s``，此时 pypdf 的 ``extract_text()`` 可能拿到乱码
（内置字体不写 ToUnicode），manifest 会用 ``cjk_font`` 字段如实记录。
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
CORPUS_DIR = REPO_ROOT / "evals" / "corpus" / "parsing"

# 系统 CJK 字体候选（按优先级）。命中第一个存在的文件即用于生成中文 PDF。
# 优先 TrueType 黑体：体积可控，子集化后单个样本仅数 KB。
CJK_FONT_CANDIDATES = (
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/msyhl.ttc",
    "C:/Windows/Fonts/simsun.ttc",
    "/usr/share/fonts/truetype/arphic/uming.ttc",
    "/System/Library/Fonts/PingFang.ttc",
)
CJK_FONT_FALLBACK = "china-s"


def _resolve_cjk_font() -> tuple[str, bool]:
    """返回 (字体名或字体文件路径, 是否为系统字体文件)。"""
    for candidate in CJK_FONT_CANDIDATES:
        if Path(candidate).exists():
            return candidate, True
    return CJK_FONT_FALLBACK, False


def _insert_cjk(page, point, text: str, font: str, fontsize: float) -> None:
    """写入中文文本；系统字体走 fontfile，内置字体走 fontname。"""
    if font.startswith("/") or Path(font).exists():
        page.insert_text(point, text, fontfile=font, fontname="cjk-corpus", fontsize=fontsize)
    else:
        page.insert_text(point, text, fontname=font, fontsize=fontsize)


def _finalize_pdf(doc) -> bytes:
    """导出 PDF 字节：先做字体子集化，再压缩清理，控制样本体积。"""
    try:
        doc.subset_fonts()
    except Exception:  # noqa: BLE001 - 老版本 PyMuPDF 无子集化能力时直接用原字体
        pass
    return doc.tobytes(garbage=4, deflate=True)


# --------------------------------------------------------------------------
# PDF：单栏多页（含标题层级）
# --------------------------------------------------------------------------

_SINGLE_COLUMN_TITLE = "产品知识库运维手册"
_SINGLE_COLUMN_CHAPTERS = [
    ("第一章 索引构建", "索引构建阶段会先解析源文件再写入向量库。"),
    ("第二章 检索与重排", "检索阶段采用向量召回加关键词召回的混合策略。"),
    ("第三章 故障排查", "故障排查优先检查依赖是否完整以及索引身份是否一致。"),
]


def build_pdf_single_column(font: str) -> bytes:
    """生成三页单栏 PDF，带 H1/H2 标题层级。"""
    import fitz

    doc = fitz.open()
    first = True
    for chapter, body in _SINGLE_COLUMN_CHAPTERS:
        page = doc.new_page(width=595, height=842)
        if first:
            _insert_cjk(page, (60, 80), _SINGLE_COLUMN_TITLE, font, 22)
            y = 140
            first = False
        else:
            y = 80
        _insert_cjk(page, (60, y), chapter, font, 16)
        _insert_cjk(page, (60, y + 40), body, font, 11)
    return _finalize_pdf(doc)


# --------------------------------------------------------------------------
# PDF：双栏（阅读顺序关键样本）
# --------------------------------------------------------------------------

_TWO_COLUMN_TITLE = "双栏排版阅读顺序样本"
_TWO_COLUMN_LEFT = [
    "左栏第一段第一行内容",
    "左栏第一段第二行内容",
    "左栏第二段第一行内容",
]
_TWO_COLUMN_RIGHT = [
    "右栏第一段第一行内容",
    "右栏第一段第二行内容",
    "右栏第二段第一行内容",
]


def build_pdf_two_column(font: str, *, interleaved: bool) -> bytes:
    """生成双栏 PDF。

    ``interleaved=False`` 按「先整左栏、再整右栏」写入内容流，模拟常见双栏
    排版（LaTeX / InDesign 导出）的内容流顺序；
    ``interleaved=True`` 按「逐行左右交替」写入，构造跨栏错乱的内容流。
    """
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    _insert_cjk(page, (60, 60), _TWO_COLUMN_TITLE, font, 18)
    y0 = 110
    if interleaved:
        for index in range(3):
            y = y0 + index * 40
            _insert_cjk(page, (60, y), _TWO_COLUMN_LEFT[index], font, 11)
            _insert_cjk(page, (320, y), _TWO_COLUMN_RIGHT[index], font, 11)
    else:
        for index, text in enumerate(_TWO_COLUMN_LEFT):
            _insert_cjk(page, (60, y0 + index * 40), text, font, 11)
        for index, text in enumerate(_TWO_COLUMN_RIGHT):
            _insert_cjk(page, (320, y0 + index * 40), text, font, 11)
    return _finalize_pdf(doc)


# --------------------------------------------------------------------------
# PDF：含表格
# --------------------------------------------------------------------------

_TABLE_TITLE = "套餐资费对照表"
_TABLE_ROWS = [
    ["套餐", "月费", "席位"],
    ["基础版", "99", "5"],
    ["专业版", "299", "20"],
    ["企业版", "899", "100"],
]


def build_pdf_table(font: str) -> bytes:
    """生成带表格的 PDF：先画框线再按单元格写文本。"""
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    _insert_cjk(page, (60, 70), _TABLE_TITLE, font, 18)
    x0, y0, col_w, row_h = 60, 130, 150, 34
    for row_index, row in enumerate(_TABLE_ROWS):
        top = y0 + row_index * row_h
        for col_index, cell in enumerate(row):
            left = x0 + col_index * col_w
            rect = fitz.Rect(left, top, left + col_w, top + row_h)
            page.draw_rect(rect, width=0.8)
            _insert_cjk(page, (left + 8, top + 22), cell, font, 11)
    return _finalize_pdf(doc)


# --------------------------------------------------------------------------
# PDF：扫描件式（图片型、无可提取文本层）
# --------------------------------------------------------------------------

_SCAN_EXPECTED_LINES = ["扫描件样例", "发票号码 12345678", "金额 1234.00 元"]


def build_pdf_scanned(font: str) -> bytes:
    """把带中文的页面渲染成位图后嵌入新 PDF，得到无文本层的扫描件样本。"""
    import fitz

    source = fitz.open()
    page = source.new_page(width=595, height=420)
    for index, line in enumerate(_SCAN_EXPECTED_LINES):
        _insert_cjk(page, (80, 100 + index * 60), line, font, 20)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    # 以 JPEG 入图，模拟扫描件常见的有损压缩位图，同时控制样本体积。
    image_bytes = pixmap.tobytes("jpeg", jpg_quality=90)

    scanned = fitz.open()
    scan_page = scanned.new_page(width=595, height=420)
    scan_page.insert_image(scan_page.rect, stream=image_bytes)
    return _finalize_pdf(scanned)


# --------------------------------------------------------------------------
# DOCX：标题层级 + 表格
# --------------------------------------------------------------------------

_DOCX_HEADINGS = ["客服系统接入说明", "接入流程", "责任矩阵"]
# 正文明避与表格单元格同词（如「联调」），保证覆盖率统计不被偶然命中污染。
_DOCX_BODY = "接入流程分为申请、验证与上线三个阶段。"
_DOCX_TABLE_ROWS = [
    ["角色", "职责", "时限"],
    ["实施", "联调", "五个工作日"],
]


def build_docx_headings_table() -> bytes:
    """生成含标题层级与表格的 docx。"""
    from docx import Document

    document = Document()
    document.add_heading(_DOCX_HEADINGS[0], level=1)
    document.add_heading(_DOCX_HEADINGS[1], level=2)
    document.add_paragraph(_DOCX_BODY)
    document.add_heading(_DOCX_HEADINGS[2], level=2)
    table = document.add_table(rows=len(_DOCX_TABLE_ROWS), cols=3)
    for row_index, row in enumerate(_DOCX_TABLE_ROWS):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# XLSX：多 sheet + 合并单元格
# --------------------------------------------------------------------------


def build_xlsx_multi_sheet() -> bytes:
    """生成含两个 sheet 与合并单元格的 xlsx。"""
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary.merge_cells("A1:B1")
    summary["A1"] = "季度汇总表"
    summary.append(["指标", "数值"])
    summary.append(["工单量", 1280])
    summary.append(["首响时长", 3.5])

    detail = workbook.create_sheet("明细")
    detail.append(["工单号", "状态"])
    detail.append(["T-1001", "已关闭"])
    detail.append(["T-1002", "处理中"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# PPTX：含备注
# --------------------------------------------------------------------------

_PPTX_SLIDES = [
    ("知识库接入培训", "本次培训覆盖解析、切分与检索三段链路。", "备注：培训材料需同步到知识库。"),
    ("验收清单", "验收清单包含覆盖率与阅读顺序两项。", "备注：验收前需完成一次全量重建。"),
]


def build_pptx_notes() -> bytes:
    """生成两页幻灯片并写入备注页文本。"""
    from pptx import Presentation

    presentation = Presentation()
    for title, body, note in _PPTX_SLIDES:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = note
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Markdown / 纯文本：代码围栏 + ASCII 盒图
# --------------------------------------------------------------------------

_ASCII_BOX = """
+-------------+        +-------------+
|  解析阶段   | -----> |  切分阶段   |
+-------------+        +-------------+
        |                     |
        v                     v
+-----------------------------+
|        检索与重排阶段        |
+-----------------------------+
""".strip()

_MARKDOWN = f"""# 部署手册

## 环境准备

部署前请确认依赖完整，脚本示例如下：

```bash
python scripts/rebuild_embedding_index.py --rebuild
```

## 链路示意

{_ASCII_BOX}

## 回滚

回滚请保留上一版索引目录，不要直接删除。
"""

_PLAIN_TEXT = f"""运维值班手册

值班期间需关注解析失败率与检索超时率两项指标。

链路示意：

{_ASCII_BOX}

异常时先查依赖，再查索引身份。
"""


def build_markdown() -> bytes:
    """返回含代码围栏与 ASCII 盒图的 Markdown 文本。"""
    return _MARKDOWN.encode("utf-8")


def build_plain_text() -> bytes:
    """返回含 ASCII 盒图的纯文本。"""
    return _PLAIN_TEXT.encode("utf-8")


# --------------------------------------------------------------------------
# 负样本
# --------------------------------------------------------------------------


def build_broken_empty() -> bytes:
    """空文件：用于验证不静默返回空知识。"""
    return b""


def build_broken_truncated() -> bytes:
    """截断的 PDF：有 PDF 头但结构损坏。"""
    return b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\ntrailer\n%%EOF\n"


def build_unsupported_binary() -> bytes:
    """不支持的二进制扩展名样本。"""
    return b"\x00\x01\x02not a document\xff\xfe"


# --------------------------------------------------------------------------
# manifest 组装
# --------------------------------------------------------------------------


def _entry(
    *,
    name: str,
    fmt: str,
    kind: str,
    expectation_mode: str,
    expected_fragments: list[str],
    expected_order: list[str] | None = None,
    expected_error: str | None = None,
    notes: str = "",
    known_gap: str | None = None,
) -> dict:
    """构造单条样本期望标注。"""
    return {
        "name": name,
        "format": fmt,
        "kind": kind,
        "expectation_mode": expectation_mode,
        "expected_fragments": expected_fragments,
        "expected_order": expected_order or [],
        "expected_error": expected_error,
        "notes": notes,
        "known_gap": known_gap,
    }


def build_manifest(cjk_font: str, font_is_system: bool) -> dict:
    """组装样本集期望标注。"""
    two_column_order = [_TWO_COLUMN_TITLE, *_TWO_COLUMN_LEFT, *_TWO_COLUMN_RIGHT]
    return {
        "version": "rag-033-parsing-corpus-v1",
        "generated_by": "scripts/generate_parse_corpus.py",
        "cjk_font": cjk_font,
        "cjk_font_is_system_font": font_is_system,
        "cjk_font_note": (
            "系统 CJK 字体带 ToUnicode，pypdf 与 PyMuPDF 均可正确提取中文"
            if font_is_system
            else "回退到 PyMuPDF 内置 china-s，pypdf 的 extract_text() 可能取到乱码"
        ),
        "samples": [
            _entry(
                name="pdf_single_column.pdf",
                fmt="pdf",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[
                    _SINGLE_COLUMN_TITLE,
                    *[chapter for chapter, _body in _SINGLE_COLUMN_CHAPTERS],
                    *[body for _chapter, body in _SINGLE_COLUMN_CHAPTERS],
                ],
                expected_order=[
                    _SINGLE_COLUMN_TITLE,
                    _SINGLE_COLUMN_CHAPTERS[0][0],
                    _SINGLE_COLUMN_CHAPTERS[0][1],
                    _SINGLE_COLUMN_CHAPTERS[1][0],
                    _SINGLE_COLUMN_CHAPTERS[1][1],
                    _SINGLE_COLUMN_CHAPTERS[2][0],
                    _SINGLE_COLUMN_CHAPTERS[2][1],
                ],
                notes="三页单栏，含 H1/H2 标题层级；顺序断言覆盖跨页顺序。",
            ),
            _entry(
                name="pdf_twocolumn.pdf",
                fmt="pdf",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[_TWO_COLUMN_TITLE, *_TWO_COLUMN_LEFT, *_TWO_COLUMN_RIGHT],
                expected_order=two_column_order,
                notes=(
                    "双栏样本，内容流按「左栏整体→右栏整体」写入；"
                    "正确阅读顺序为左栏全部行先于右栏全部行。"
                ),
            ),
            _entry(
                name="pdf_twocolumn_interleaved.pdf",
                fmt="pdf",
                kind="positive",
                expectation_mode="known_limitation",
                expected_fragments=[_TWO_COLUMN_TITLE, *_TWO_COLUMN_LEFT, *_TWO_COLUMN_RIGHT],
                expected_order=two_column_order,
                notes=(
                    "与 pdf_twocolumn.pdf 文本完全相同，但内容流按「逐行左右交替」写入。"
                    "当前解析器沿用内容流顺序、未做几何分栏排序，本样本用于量化跨栏错乱。"
                ),
                known_gap="解析器未做列聚类排序，逐行交替写入的双栏 PDF 阅读顺序会跨栏错乱。",
            ),
            _entry(
                name="pdf_table.pdf",
                fmt="pdf",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[
                    _TABLE_TITLE,
                    *[cell for row in _TABLE_ROWS for cell in row],
                ],
                notes="带框线表格的 PDF；覆盖率按单元格取值逐项命中。",
            ),
            _entry(
                name="pdf_scanned.pdf",
                fmt="pdf",
                kind="ocr",
                expectation_mode="must_pass",
                expected_fragments=_SCAN_EXPECTED_LINES,
                notes=(
                    "图片型 PDF（中文页面渲染成位图后嵌入），无文本层，必然走 OCR 路径。"
                    "本机无 tesseract 二进制，真实 provider 必然报依赖缺失。"
                ),
            ),
            _entry(
                name="docx_headings_table.docx",
                fmt="docx",
                kind="positive",
                expectation_mode="known_limitation",
                expected_fragments=[
                    *_DOCX_HEADINGS,
                    _DOCX_BODY,
                    *[cell for row in _DOCX_TABLE_ROWS for cell in row],
                ],
                notes="含 H1/H2 标题层级与 2x3 表格。",
                known_gap=(
                    "DocxDocumentParser 只遍历 document.paragraphs，表格单元格不进入解析结果，"
                    "故表格片段必然未命中。"
                ),
            ),
            _entry(
                name="xlsx_multi_sheet.xlsx",
                fmt="xlsx",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[
                    "# Sheet: 汇总",
                    "季度汇总表",
                    "指标",
                    "数值",
                    "工单量",
                    "1280",
                    "首响时长",
                    "3.5",
                    "# Sheet: 明细",
                    "工单号",
                    "状态",
                    "T-1001",
                    "已关闭",
                    "T-1002",
                    "处理中",
                ],
                notes="两个 sheet，A1:B1 为合并单元格；合并单元格取值只出现在左上角。",
            ),
            _entry(
                name="pptx_notes.pptx",
                fmt="pptx",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[title for title, _body, _note in _PPTX_SLIDES]
                + [body for _title, body, _note in _PPTX_SLIDES]
                + [note for _title, _body, note in _PPTX_SLIDES],
                notes="两页幻灯片，每页带备注页文本。",
            ),
            _entry(
                name="md_code_fence.md",
                fmt="md",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[
                    "# 部署手册",
                    "## 环境准备",
                    "```bash",
                    "python scripts/rebuild_embedding_index.py --rebuild",
                    "```",
                    "+---",
                    "解析阶段",
                    "切分阶段",
                    "检索与重排阶段",
                    "## 回滚",
                ],
                notes="含 bash 代码围栏与 ASCII 盒图；盒图完整性呼应 RAG-042 的展示诉求。",
            ),
            _entry(
                name="txt_plain.txt",
                fmt="txt",
                kind="positive",
                expectation_mode="must_pass",
                expected_fragments=[
                    "运维值班手册",
                    "解析失败率",
                    "检索超时率",
                    "链路示意",
                    "+-------------+",
                    "解析阶段",
                    "检索与重排阶段",
                    "异常时先查依赖",
                ],
                notes="纯文本，含 ASCII 盒图。",
            ),
            _entry(
                name="broken_empty.pdf",
                fmt="pdf",
                kind="negative",
                expectation_mode="must_fail",
                expected_fragments=[],
                expected_error="DocumentParseError",
                notes="零字节 PDF，验证不静默返回空知识。",
            ),
            _entry(
                name="broken_truncated.pdf",
                fmt="pdf",
                kind="negative",
                expectation_mode="must_fail",
                expected_fragments=[],
                expected_error="DocumentParseError",
                notes="有 PDF 头但结构损坏的截断文件。",
            ),
            _entry(
                name="unsupported.bin",
                fmt="bin",
                kind="negative",
                expectation_mode="must_fail",
                expected_fragments=[],
                expected_error="UnsupportedDocumentTypeError",
                notes="不支持的扩展名，注册表应明确报错。",
            ),
        ],
    }


def generate() -> dict:
    """生成全部样本文件与 manifest，返回 manifest。"""
    import fitz  # noqa: F401  确认 PyMuPDF 可用

    font, font_is_system = _resolve_cjk_font()
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)

    builders = {
        "pdf_single_column.pdf": lambda: build_pdf_single_column(font),
        "pdf_twocolumn.pdf": lambda: build_pdf_two_column(font, interleaved=False),
        "pdf_twocolumn_interleaved.pdf": lambda: build_pdf_two_column(font, interleaved=True),
        "pdf_table.pdf": lambda: build_pdf_table(font),
        "pdf_scanned.pdf": lambda: build_pdf_scanned(font),
        "docx_headings_table.docx": build_docx_headings_table,
        "xlsx_multi_sheet.xlsx": build_xlsx_multi_sheet,
        "pptx_notes.pptx": build_pptx_notes,
        "md_code_fence.md": build_markdown,
        "txt_plain.txt": build_plain_text,
        "broken_empty.pdf": build_broken_empty,
        "broken_truncated.pdf": build_broken_truncated,
        "unsupported.bin": build_unsupported_binary,
    }

    for name, builder in builders.items():
        (CORPUS_DIR / name).write_bytes(builder())

    manifest = build_manifest(font, font_is_system)
    (CORPUS_DIR / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return manifest


def main() -> int:
    """命令行入口：生成样本集并打印清单。"""
    manifest = generate()
    print(f"样本集已生成：{CORPUS_DIR}")
    print(f"CJK 字体：{manifest['cjk_font']}（系统字体={manifest['cjk_font_is_system_font']}）")
    for sample in manifest["samples"]:
        path = CORPUS_DIR / sample["name"]
        print(f"  - {sample['name']:<34} {path.stat().st_size:>8} B  {sample['kind']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
