"""Independent request-level attachment boundaries and persistence checks."""

import base64
import io
import json

import pytest

from tests.test_chat_attachments import _counts
from tests.test_chat_attachments import attachment_client as _attachment_client

attachment_client = _attachment_client


def item(name, content):
    return {"name": name, "content_base64": base64.b64encode(content).decode("ascii")}


@pytest.mark.parametrize("attachment, status", [
    (item("view.jsx", b"export default 1"), 400),
    (item("plain.txt", b"not an allowed extension"), 400),
    (item("huge.py", b"x" * (1048576 + 1)), 413),
    ({"name": "source.py", "content_base64": "%%%%"}, 400),
    (item("source.py", b"\xff\xfeinvalid utf8"), 400),
    (item("photo.png", b"not a png"), 400),
    (item("photo.jpg", b"not a jpeg"), 400),
    (item("broken.pdf", b"not a pdf"), 400),
])
def test_invalid_batch_has_no_message_knowledge_or_model_side_effect(attachment_client, attachment, status):
    client, _user, model = attachment_client
    before = _counts()
    response = client.post("/api/v1/chat", json={
        "message": "review this batch",
        "attachments": [item("valid.py", b"VALID_FIRST_ITEM"), attachment],
    })
    assert response.status_code == status
    if status == 413:
        error = response.json()["error"]
        assert error["code"] == "file_too_large"
        assert "1048576" in error["message"]
    assert _counts() == before
    assert model.requests == []
    assert "Traceback" not in response.text
    assert "content_base64" not in response.text


def test_all_source_extensions_bind_to_current_message_without_a_count_limit(attachment_client):
    client, _user, model = attachment_client
    extensions = "go py css html htm java js json xml yaml yml c h cpp cc cxx hpp vue ts".split()
    before = _counts()
    files = [item(f"sample.{ext}", f"REVIEW_SOURCE_{ext}_BODY".encode()) for ext in extensions]
    response = client.post("/api/v1/chat", json={"message": "read all sources", "attachments": files})
    assert response.status_code == 200
    text = json.dumps([[message.model_dump() for message in request] for request in model.requests])
    for ext in extensions:
        assert f"REVIEW_SOURCE_{ext}_BODY" in text
    detail = client.get(f"/api/v1/chat/conversations/{response.json()['conversation_id']}")
    assert detail.status_code == 200
    user_messages = [message for message in detail.json()["messages"] if message["role"] == "user"]
    assert len(user_messages) == 1
    assert len(user_messages[0]["attachments"]) == len(extensions)
    assert all(set(summary) == {"name", "size", "mime_type"} for summary in user_messages[0]["attachments"])
    assert "content_base64" not in detail.text
    assert _counts()[1:] == before[1:]


def test_attachment_only_message_is_supported_but_empty_message_is_rejected(attachment_client):
    client, _user, model = attachment_client
    response = client.post("/api/v1/chat", json={
        "message": "", "attachments": [item("question.md", b"REVIEW_ATTACHMENT_ONLY")],
    })
    assert response.status_code == 200
    assert any("REVIEW_ATTACHMENT_ONLY" in str(message.content)
               for request in model.requests for message in request)
    before = _counts()
    calls = len(model.requests)
    response = client.post("/api/v1/chat", json={"message": "", "attachments": []})
    assert response.status_code == 400
    assert _counts() == before
    assert len(model.requests) == calls


def test_invalid_stream_attachment_is_rejected_before_run_creation(attachment_client):
    from app.services.active_runs import active_runs

    client, _user, model = attachment_client
    before, runs = _counts(), set(active_runs.runs)
    response = client.post("/api/v1/chat/stream", json={
        "message": "stream invalid", "attachments": [item("bad.tsx", b"bad")],
    })
    assert response.status_code == 400
    assert _counts() == before
    assert set(active_runs.runs) == runs
    assert model.requests == []


def test_real_modern_documents_reach_model_without_knowledge_ingestion(attachment_client):
    import pymupdf
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    document = Document()
    document.add_paragraph("REVIEW_DOCX_BODY")
    word = io.BytesIO()
    document.save(word)
    workbook = Workbook()
    workbook.active["A1"] = "REVIEW_XLSX_BODY"
    sheet = io.BytesIO()
    workbook.save(sheet)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "REVIEW_PPTX_BODY"
    slides = io.BytesIO()
    presentation.save(slides)
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "REVIEW_PDF_BODY")
        pdf_bytes = pdf.tobytes()
    files = [item("sample.docx", word.getvalue()), item("sample.xlsx", sheet.getvalue()),
             item("sample.pptx", slides.getvalue()), item("sample.pdf", pdf_bytes),
             item("sample.md", b"REVIEW_MARKDOWN_BODY")]
    client, _user, model = attachment_client
    before = _counts()
    response = client.post("/api/v1/chat", json={"message": "read documents", "attachments": files})
    assert response.status_code == 200
    payload = json.dumps([[message.model_dump() for message in request] for request in model.requests])
    for marker in ("REVIEW_DOCX_BODY", "REVIEW_XLSX_BODY", "REVIEW_PPTX_BODY", "REVIEW_PDF_BODY",
                   "REVIEW_MARKDOWN_BODY"):
        assert marker in payload
    assert _counts()[1:] == before[1:]


def test_unverified_image_provider_rejects_before_any_message_write(attachment_client):
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    client, _user, model = attachment_client
    before = _counts()
    response = client.post("/api/v1/chat", json={
        "message": "read picture", "attachments": [item("photo.png", image.getvalue())],
    })
    assert response.status_code == 400
    assert _counts() == before
    assert model.requests == []


def _vision_recorder():
    from app.llm.mock import MockLLMProvider
    from app.llm.openai_compatible import OpenAICompatibleProvider

    class VisionRecorder(OpenAICompatibleProvider):
        def __init__(self):
            super().__init__("https://api.openai.com/v1", "offline-review-placeholder", "gpt-4o-mini")
            self.mock = MockLLMProvider()
            self.payloads = []

        async def chat(self, messages, options=None):
            self._guard(messages, options)
            self.payloads.append([message.model_dump(mode="json") for message in messages])
            return await self.mock.chat(messages, options)

        async def stream_chat(self, messages, options=None):
            yield await self.chat(messages, options)

    return VisionRecorder()


def test_valid_image_reaches_guarded_model_and_history_exposes_only_metadata(attachment_client):
    from PIL import Image

    from app.llm.factory import set_llm_provider_override

    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    encoded = base64.b64encode(image.getvalue()).decode("ascii")
    client, _user, _model = attachment_client
    recorder = _vision_recorder()
    set_llm_provider_override(recorder)
    before = _counts()
    response = client.post("/api/v1/chat", json={
        "message": "read picture", "attachments": [{"name": "photo.png", "content_base64": encoded}],
    })
    assert response.status_code == 200
    assert any(any(isinstance(message["content"], list) and any(
        part.get("type") == "image_url" and part["image_url"] == {
            "url": "data:image/png;base64," + encoded, "detail": "low",
        } for part in message["content"]) for message in request) for request in recorder.payloads)
    detail = client.get(f"/api/v1/chat/conversations/{response.json()['conversation_id']}")
    assert detail.status_code == 200
    assert encoded not in detail.text
    summaries = [message["attachments"] for message in detail.json()["messages"] if message["role"] == "user"]
    assert summaries == [[{"name": "photo.png", "size": len(image.getvalue()), "mime_type": "image/png"}]]
    assert _counts()[1:] == before[1:]


def test_exact_one_mib_is_accepted_by_file_validation():
    from app.services.chat_attachments import prepare_attachments

    prepared = prepare_attachments([item("boundary.py", b"x" * 1048576)])
    assert len(prepared) == 1 and prepared[0].size == 1048576
    assert len(prepared[0].text) == 1048576


@pytest.mark.parametrize("streaming", [False, True])
def test_resume_reconstructs_original_image_and_text_without_extra_user_message(
    attachment_client, monkeypatch, streaming,
):
    from PIL import Image
    from sqlmodel import Session, col, select

    from app.agents.pipeline import AgentState
    from app.agents.state import StateManager
    from app.core.config import settings
    from app.core.database import engine
    from app.llm.factory import set_llm_provider_override
    from app.models.conversation import Conversation, Message
    from app.models.user import Tenant
    from app.services.agent_state_service import AgentStateService
    from app.services.chat_attachments import prepare_attachments, serialize_attachments

    client, user, _model = attachment_client
    monkeypatch.setattr(settings, "AGENT_STATE_PERSIST_ENABLED", True)
    monkeypatch.setattr(settings, "AGENT_ORCHESTRATION", "self")
    recorder = _vision_recorder()
    set_llm_provider_override(recorder)
    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    original = [item("original.png", image.getvalue()), item("original.md", b"RESUME_ORIGINAL_BODY")]
    with Session(engine) as session:
        tenant = session.get(Tenant, user.tenant_id)
        tenant.message_limit = 1
        session.add(tenant)
        conversation = Conversation(tenant_id=user.tenant_id, user_id=user.id)
        session.add(conversation)
        session.commit()
        message = Message(conversation_id=conversation.id, role="user", content="original question",
                          attachments=serialize_attachments(prepare_attachments(original)))
        session.add(message)
        session.commit()
        conversation_id, message_id = conversation.id, message.id
        state = AgentState(user_input=message.content, needs_full_pipeline=False)
        AgentStateService(session, tenant_id=user.tenant_id, user_id=user.id).save_checkpoint(
            conversation_id=conversation_id, message_id=message_id, tenant_id=user.tenant_id, user_id=user.id,
            stage="action", round_no=0, tool_budget_exhausted=False, state_json=StateManager(state).to_json(),
            status="running", error=None,
        )
    response = client.post("/api/v1/chat/stream" if streaming else "/api/v1/chat", json={
        "message": "original question", "conversation_id": conversation_id, "resume_message_id": message_id,
    })
    assert response.status_code == 200
    payload = json.dumps(recorder.payloads)
    assert "RESUME_ORIGINAL_BODY" in payload
    assert "data:image/png;base64," + original[0]["content_base64"] in payload
    with Session(engine) as session:
        messages = session.exec(select(Message).where(col(Message.conversation_id) == conversation_id)).all()
        assert len([message for message in messages if message.role == "user"]) == 1
        assert any(message.role == "assistant" and message.status == "complete" for message in messages)
    before, calls = _counts(), len(recorder.payloads)
    replacement = client.post("/api/v1/chat", json={
        "message": "original question", "conversation_id": conversation_id, "resume_message_id": message_id,
        "attachments": [item("replacement.md", b"REPLACEMENT_BODY")],
    })
    assert replacement.status_code == 400
    assert _counts() == before and len(recorder.payloads) == calls


@pytest.mark.parametrize("other_tenant", [False, True])
def test_other_actor_cannot_read_or_resume_attachment_message(attachment_client, persist_identity, other_tenant):
    import secrets

    from app.api.deps import get_current_user
    from app.main import app
    from app.models.user import User

    client, user, model = attachment_client
    created = client.post("/api/v1/chat", json={
        "message": "original request", "attachments": [item("private.md", b"PRIVATE_ATTACHMENT_BODY")],
    })
    assert created.status_code == 200
    body = created.json()
    identity = "other-attachment-" + secrets.token_hex(8)
    other = User(id=identity, tenant_id=identity if other_tenant else user.tenant_id,
                 username=identity, hashed_password="", role="member")
    persist_identity(other)
    app.dependency_overrides[get_current_user] = lambda: other
    before, calls = _counts(), len(model.requests)
    detail = client.get(f"/api/v1/chat/conversations/{body['conversation_id']}")
    assert detail.status_code == 404 and "PRIVATE_ATTACHMENT_BODY" not in detail.text
    resumed = client.post("/api/v1/chat", json={
        "message": "original request", "conversation_id": body["conversation_id"],
        "resume_message_id": body["message_id"],
    })
    assert resumed.status_code == 404
    assert _counts() == before and len(model.requests) == calls
