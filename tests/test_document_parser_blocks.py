import pytest


def test_parsed_document_blocks_default_empty() -> None:
    from app.rag.document_parsers.base import ParsedDocument

    doc = ParsedDocument(
        text="正文",
        title="标题",
        source="a.docx",
        extension="docx",
        content_type=None,
    )
    assert doc.blocks == []


def test_service_preserves_blocks() -> None:
    from app.rag.document_parsers.base import DocumentParser, ParsedBlock, ParsedDocument
    from app.rag.document_parsers.registry import DocumentParserRegistry
    from app.rag.document_parsers.service import DocumentParserService

    class BlockParser(DocumentParser):
        def supports(self, filename: str, content_type: str | None) -> bool:
            return filename.endswith(".txt")

        def extract(
            self, file_bytes: bytes, filename: str, content_type: str | None
        ) -> ParsedDocument:
            return ParsedDocument(
                text="第一行\n第二行",
                title="块测试",
                source=filename,
                extension="txt",
                content_type=content_type,
                metadata={"parser_name": "fake"},
                blocks=[
                    ParsedBlock(type="text", text="第一行", order=0),
                    ParsedBlock(type="text", text="第二行", order=1),
                ],
            )

    service = DocumentParserService(DocumentParserRegistry([BlockParser()]))
    parsed = service.parse_upload(b"hello", "demo.txt", "text/plain")
    assert len(parsed.blocks) == 2
    assert [block.text for block in parsed.blocks] == ["第一行", "第二行"]


def test_docx_parser_outputs_blocks(tmp_path) -> None:
    from docx import Document

    from app.rag.document_parsers.service import parse_uploaded_document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    document.save(path)

    parsed = parse_uploaded_document(
        path.read_bytes(),
        path.name,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert len(parsed.blocks) == 2
    assert [block.text for block in parsed.blocks] == ["第一段", "第二段"]
    assert all(block.type == "paragraph" for block in parsed.blocks)
    assert [block.metadata["reading_order"] for block in parsed.blocks] == [1, 2]
    assert all(block.metadata["layout_role"] == "body" for block in parsed.blocks)


def test_xlsx_parser_outputs_sheet_blocks(tmp_path) -> None:
    from openpyxl import Workbook

    from app.rag.document_parsers.service import parse_uploaded_document

    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工"
    sheet["A1"] = "姓名"
    sheet["B1"] = "部门"
    sheet["A2"] = "Alice"
    sheet["B2"] = "销售"
    workbook.save(path)

    parsed = parse_uploaded_document(
        path.read_bytes(),
        path.name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert parsed.blocks
    assert parsed.blocks[0].type == "sheet_heading"
    assert parsed.blocks[0].section_path == ["员工"]
    assert any(block.type == "sheet_row" and "Alice" in block.text for block in parsed.blocks)


def test_pdf_parser_outputs_page_blocks(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.pdf import PdfDocumentParser

    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        is_encrypted = False
        pages = [FakePage("第一页内容"), FakePage("第二页内容")]

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    parsed = PdfDocumentParser().extract(b"%PDF-1.4", "text.pdf", "application/pdf")
    assert [block.page for block in parsed.blocks] == [1, 2]
    assert [block.text for block in parsed.blocks] == ["第一页内容", "第二页内容"]
    assert [block.metadata["reading_order"] for block in parsed.blocks] == [1, 1]
    assert all(block.metadata["layout_role"] == "body" for block in parsed.blocks)


def test_pdf_parser_outputs_native_text_block_bboxes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.rag.document_parsers.pdf import PdfDocumentParser

    class FakePage:
        def extract_text(self) -> str:
            return "标题\n正文内容"

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class FakeFitzPage:
        def get_text(self, mode: str):
            assert mode == "blocks"
            return [
                (10.0, 20.0, 200.0, 40.0, "标题\n", 0, 0),
                (10.0, 60.0, 210.0, 90.0, "正文内容\n", 1, 0),
            ]

    class FakeFitzDocument:
        def __iter__(self):
            return iter([FakeFitzPage()])

    class FakeFitz:
        @staticmethod
        def open(stream, filetype):
            return FakeFitzDocument()

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    monkeypatch.setattr("app.rag.document_parsers.pdf.fitz", FakeFitz())
    parsed = PdfDocumentParser().extract(b"%PDF-1.4", "text.pdf", "application/pdf")

    assert [block.text for block in parsed.blocks] == ["标题", "正文内容"]
    assert parsed.blocks[0].bbox == (10.0, 20.0, 200.0, 40.0)
    assert parsed.blocks[1].bbox == (10.0, 60.0, 210.0, 90.0)
    assert [block.metadata["reading_order"] for block in parsed.blocks] == [1, 2]
