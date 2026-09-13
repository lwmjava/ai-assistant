import io
import zipfile
from pathlib import Path

import pytest


def test_registry_rejects_unsupported_extension() -> None:
    from app.rag.document_parsers.base import UnsupportedDocumentTypeError
    from app.rag.document_parsers.service import parse_uploaded_document

    with pytest.raises(UnsupportedDocumentTypeError):
        parse_uploaded_document(b"binary", "data.bin", "application/octet-stream")


def test_service_rejects_empty_text() -> None:
    from app.rag.document_parsers.base import (
        DocumentParser,
        DocumentTextEmptyError,
        ParsedDocument,
    )
    from app.rag.document_parsers.registry import DocumentParserRegistry
    from app.rag.document_parsers.service import DocumentParserService

    class EmptyParser(DocumentParser):
        def supports(self, filename: str, content_type: str | None) -> bool:
            return filename.endswith(".txt")

        def extract(
            self, file_bytes: bytes, filename: str, content_type: str | None
        ) -> ParsedDocument:
            return ParsedDocument(
                text="   ",
                title="a",
                source=filename,
                extension="txt",
                content_type=content_type,
                metadata={},
            )

    service = DocumentParserService(DocumentParserRegistry([EmptyParser()]))
    with pytest.raises(DocumentTextEmptyError):
        service.parse_upload(b"abc", "a.txt", "text/plain")


def test_parse_json_file() -> None:
    from app.rag.document_parsers.service import parse_uploaded_document

    parsed = parse_uploaded_document(
        b'{"k": "v"}', "doc.json", "application/json"
    )
    assert parsed.extension == "json"
    assert '"k": "v"' in parsed.text


def test_parse_csv_with_gbk_fallback() -> None:
    from app.rag.document_parsers.service import parse_uploaded_document

    parsed = parse_uploaded_document(
        "姓名,部门".encode("gb18030"), "staff.csv", "text/csv"
    )
    assert "姓名" in parsed.text


def test_docx_parser_extracts_text(tmp_path: Path) -> None:
    from docx import Document

    from app.rag.document_parsers.service import parse_uploaded_document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("企业知识库文档")
    document.save(path)

    parsed = parse_uploaded_document(
        path.read_bytes(),
        path.name,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "企业知识库文档" in parsed.text


def test_xlsx_parser_extracts_text(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from app.rag.document_parsers.service import parse_uploaded_document

    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工"
    sheet["A1"] = "姓名"
    sheet["B1"] = "部门"
    sheet["A2"] = "Alice"
    sheet["B2"] = "销售"
    workbook.save(path)

    parsed = parse_uploaded_document(
        path.read_bytes(),
        path.name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert "Alice" in parsed.text
    assert "销售" in parsed.text


def test_pptx_parser_extracts_text(tmp_path: Path) -> None:
    from pptx import Presentation

    from app.rag.document_parsers.service import parse_uploaded_document

    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度汇报"
    slide.placeholders[1].text = "销售额增长 20%"
    presentation.save(path)

    parsed = parse_uploaded_document(
        path.read_bytes(),
        path.name,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "季度汇报" in parsed.text
    assert "销售额增长 20%" in parsed.text


def test_pptx_parser_extracts_group_table_and_notes(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.office import PptxDocumentParser

    class FakeRun:
        def __init__(self, text: str) -> None:
            self.text = text

    class FakeParagraph:
        def __init__(self, *texts: str) -> None:
            self.runs = [FakeRun(text) for text in texts]

    class FakeTextFrame:
        def __init__(self, *paragraphs: tuple[str, ...]) -> None:
            self.paragraphs = [FakeParagraph(*paragraph) for paragraph in paragraphs]

    class FakeCell:
        def __init__(self, text: str) -> None:
            self.text_frame = FakeTextFrame((text,))

    class FakeRow:
        def __init__(self, *texts: str) -> None:
            self.cells = [FakeCell(text) for text in texts]

    class FakeTable:
        def __init__(self) -> None:
            self.rows = [FakeRow("表头", "值"), FakeRow("销售额", "100")]

    class FakeShape:
        def __init__(
            self,
            *,
            text_frame: FakeTextFrame | None = None,
            table: FakeTable | None = None,
            shapes: list["FakeShape"] | None = None,
            text: str = "",
            left: int = 0,
            top: int = 0,
            width: int = 100,
            height: int = 50,
            name: str = "",
        ) -> None:
            self.has_text_frame = text_frame is not None
            self.text_frame = text_frame
            self.has_table = table is not None
            self.table = table
            self.shapes = shapes
            self.text = text
            self.left = left
            self.top = top
            self.width = width
            self.height = height
            self.name = name

    class FakeNotesSlide:
        def __init__(self) -> None:
            self.shapes = [FakeShape(text_frame=FakeTextFrame(("备注信息",)), top=400, name="notes")]

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = [
                FakeShape(text_frame=FakeTextFrame(("标题",)), top=10, height=40, name="title"),
                FakeShape(table=FakeTable(), top=100, height=120, name="table"),
                FakeShape(
                    shapes=[FakeShape(text_frame=FakeTextFrame(("分组文本",)), left=300, top=250)],
                    top=250,
                    name="group",
                ),
            ]
            self.has_notes_slide = True
            self.notes_slide = FakeNotesSlide()

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)
    parsed = PptxDocumentParser().extract(
        b"fake-pptx", "complex.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "标题" in parsed.text
    assert "表头" in parsed.text
    assert "销售额" in parsed.text
    assert "分组文本" in parsed.text
    assert "备注信息" in parsed.text
    title_block = next(block for block in parsed.blocks if block.text == "标题")
    assert title_block.bbox == (0.0, 10.0, 100.0, 50.0)
    assert title_block.metadata["layout_role"] == "title"
    table_block = next(block for block in parsed.blocks if block.text == "表头")
    assert table_block.metadata["layout_role"] == "table_header"
    note_block = next(block for block in parsed.blocks if block.text == "备注信息")
    assert note_block.metadata["layout_role"] == "note"


def test_pptx_parser_extracts_chart_text(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.office import PptxDocumentParser

    class FakeTextFrame:
        def __init__(self, text: str) -> None:
            self.paragraphs = [type("Paragraph", (), {"runs": [type("Run", (), {"text": text})()]})()]

    class FakeChartTitle:
        def __init__(self, text: str) -> None:
            self.has_text_frame = True
            self.text_frame = FakeTextFrame(text)

    class FakeAxisTitle:
        def __init__(self, text: str) -> None:
            self.has_text_frame = True
            self.text_frame = FakeTextFrame(text)

    class FakeAxis:
        def __init__(self, text: str) -> None:
            self.has_title = True
            self.axis_title = FakeAxisTitle(text)

    class FakeSeries:
        def __init__(self, name: str) -> None:
            self.name = name

    class FakeCategory:
        def __init__(self, label: str) -> None:
            self.label = label

    class FakePlot:
        def __init__(self) -> None:
            self.series = [FakeSeries("华东区"), FakeSeries("华南区")]
            self.categories = [FakeCategory("Q1"), FakeCategory("Q2")]

    class FakeChart:
        def __init__(self) -> None:
            self.has_title = True
            self.chart_title = FakeChartTitle("销售趋势")
            self.category_axis = FakeAxis("季度")
            self.value_axis = FakeAxis("销售额")
            self.plots = [FakePlot()]

    class FakeShape:
        def __init__(self) -> None:
            self.has_text_frame = False
            self.has_table = False
            self.has_chart = True
            self.chart = FakeChart()
            self.shapes = None
            self.text = ""
            self.left = 120
            self.top = 80
            self.width = 320
            self.height = 180
            self.name = "chart"

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = [FakeShape()]
            self.has_notes_slide = False

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)
    parsed = PptxDocumentParser().extract(
        b"fake-chart-pptx",
        "chart.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "销售趋势" in parsed.text
    assert "季度" in parsed.text
    assert "销售额" in parsed.text
    assert "华东区" in parsed.text
    assert "Q1" in parsed.text
    chart_title_block = next(block for block in parsed.blocks if block.text == "销售趋势")
    assert chart_title_block.bbox == (120.0, 80.0, 440.0, 260.0)
    assert chart_title_block.metadata["layout_role"] == "chart_title"
    category_block = next(block for block in parsed.blocks if block.text == "Q1")
    assert category_block.metadata["layout_role"] == "chart_category"


def test_pptx_parser_extracts_chart_legend_and_data_labels(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.rag.document_parsers.office import PptxDocumentParser

    class FakeTextFrame:
        def __init__(self, text: str) -> None:
            self.paragraphs = [type("Paragraph", (), {"runs": [type("Run", (), {"text": text})()]})()]

    class FakeLegendEntry:
        def __init__(self, text: str) -> None:
            self.text_frame = FakeTextFrame(text)

    class FakeLegend:
        def __init__(self) -> None:
            self.entries = [FakeLegendEntry("图例-华东"), FakeLegendEntry("图例-华南")]

    class FakeDataLabel:
        def __init__(self, text: str) -> None:
            self.has_text_frame = True
            self.text_frame = FakeTextFrame(text)

    class FakePoint:
        def __init__(self, text: str) -> None:
            self.data_label = FakeDataLabel(text)

    class FakeSeries:
        def __init__(self, name: str) -> None:
            self.name = name
            self.points = [FakePoint(f"{name}-数据标签")]

    class FakePlot:
        def __init__(self) -> None:
            self.series = [FakeSeries("华东区")]
            self.categories = []

    class FakeChart:
        def __init__(self) -> None:
            self.has_title = False
            self.has_legend = True
            self.legend = FakeLegend()
            self.plots = [FakePlot()]

    class FakeShape:
        def __init__(self) -> None:
            self.has_text_frame = False
            self.has_table = False
            self.has_chart = True
            self.chart = FakeChart()
            self.shapes = None
            self.text = ""

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = [FakeShape()]
            self.has_notes_slide = False

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)
    parsed = PptxDocumentParser().extract(
        b"fake-chart-pptx",
        "chart_legend.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "图例-华东" in parsed.text
    assert "图例-华南" in parsed.text
    assert "华东区-数据标签" in parsed.text


def test_pptx_parser_falls_back_to_ooxml_text() -> None:
    from app.rag.document_parsers.office import PptxDocumentParser

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "ppt/slides/slide1.xml",
            """
            <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <p:cSld>
                    <p:spTree>
                        <p:sp>
                            <p:txBody>
                                <a:p><a:r><a:t>XML 回退文本</a:t></a:r></a:p>
                            </p:txBody>
                        </p:sp>
                    </p:spTree>
                </p:cSld>
            </p:sld>
            """,
        )
        archive.writestr(
            "ppt/notesSlides/notesSlide1.xml",
            """
            <p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <p:cSld>
                    <p:spTree>
                        <p:sp>
                            <p:txBody>
                                <a:p><a:r><a:t>XML 备注文本</a:t></a:r></a:p>
                            </p:txBody>
                        </p:sp>
                    </p:spTree>
                </p:cSld>
            </p:notes>
            """,
        )

    parsed = PptxDocumentParser().extract(
        buffer.getvalue(),
        "xml_fallback.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "XML 回退文本" in parsed.text
    assert "XML 备注文本" in parsed.text


def test_pptx_parser_falls_back_to_smartart_ooxml_text() -> None:
    from app.rag.document_parsers.office import PptxDocumentParser

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "ppt/diagrams/data1.xml",
            """
            <dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
                           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <dgm:ptLst>
                    <dgm:pt modelId="0">
                        <dgm:t>
                            <a:p><a:r><a:t>线索收集</a:t></a:r></a:p>
                        </dgm:t>
                    </dgm:pt>
                    <dgm:pt modelId="1">
                        <dgm:t>
                            <a:p><a:r><a:t>线索分发</a:t></a:r></a:p>
                        </dgm:t>
                    </dgm:pt>
                </dgm:ptLst>
            </dgm:dataModel>
            """,
        )

    parsed = PptxDocumentParser().extract(
        buffer.getvalue(),
        "smartart.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "线索收集" in parsed.text
    assert "线索分发" in parsed.text


def test_pptx_parser_uses_pdf_ocr_fallback_when_text_empty(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.rag.document_parsers.office import PptxDocumentParser
    from app.rag.ocr.base import OcrPageBlock, OcrResult

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = []
            self.has_notes_slide = False

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    class FakeProvider:
        def extract_pdf_text(self, file_bytes: bytes) -> OcrResult:
            assert file_bytes == b"%PDF-1.4 fake"
            return OcrResult(
                text="OCR 截图页文本",
                provider="tesseract",
                used_ocr=True,
                blocks=[
                    OcrPageBlock(
                        page=1,
                        text="OCR 截图页文本",
                        bbox=(0.0, 0.0, 1280.0, 720.0),
                        layout_role="body",
                    )
                ],
            )

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)
    monkeypatch.setattr(
        "app.rag.document_parsers.office._convert_office_bytes_with_libreoffice",
        lambda file_bytes, source_ext, target_ext: b"%PDF-1.4 fake",
    )
    parsed = PptxDocumentParser(ocr_provider=FakeProvider()).extract(
        b"fake-pptx",
        "scan_deck.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert parsed.text == "OCR 截图页文本"
    assert parsed.metadata["used_ocr"] is True
    assert parsed.metadata["ocr_provider"] == "tesseract"
    assert parsed.metadata["parser_name"] == "pptx_ocr"
    assert parsed.blocks[0].bbox == (0.0, 0.0, 1280.0, 720.0)
    assert parsed.blocks[0].metadata["layout_role"] == "body"


def test_pdf_parser_reports_ocr_required(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.base import DocumentOcrRequiredError
    from app.rag.document_parsers.pdf import PdfDocumentParser

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    with pytest.raises(DocumentOcrRequiredError):
        PdfDocumentParser().extract(b"%PDF", "scan.pdf", "application/pdf")


def test_pdf_parser_uses_text_layer_without_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.pdf import PdfDocumentParser

    class FakePage:
        def extract_text(self) -> str:
            return "PDF 自带文本层"

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    parsed = PdfDocumentParser().extract(b"%PDF-1.4", "text.pdf", "application/pdf")
    assert parsed.text == "PDF 自带文本层"
    assert parsed.metadata["used_ocr"] is False


def test_pdf_parser_uses_ocr_when_text_layer_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrPageBlock, OcrResult

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class FakeProvider:
        def extract_pdf_text(self, file_bytes: bytes) -> OcrResult:
            return OcrResult(
                text="OCR 提取文本",
                provider="tesseract",
                used_ocr=True,
                blocks=[
                    OcrPageBlock(
                        page=1,
                        text="OCR 提取文本",
                        bbox=(0.0, 0.0, 1440.0, 1920.0),
                    )
                ],
            )

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    parsed = PdfDocumentParser(ocr_provider=FakeProvider()).extract(
        b"%PDF-1.4", "scan.pdf", "application/pdf"
    )
    assert parsed.text == "OCR 提取文本"
    assert parsed.metadata["used_ocr"] is True
    assert parsed.metadata["parser_name"] == "pdf_ocr"
    assert parsed.blocks[0].bbox == (0.0, 0.0, 1440.0, 1920.0)


def test_pdf_parser_preserves_cloud_ocr_provider_name(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrResult

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class CloudProvider:
        def extract_pdf_text(self, file_bytes: bytes) -> OcrResult:
            return OcrResult(
                text="云 OCR 提取文本",
                provider="cloud_openai_vision",
                used_ocr=True,
            )

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    parsed = PdfDocumentParser(ocr_provider=CloudProvider()).extract(
        b"%PDF-1.4", "scan.pdf", "application/pdf"
    )
    assert parsed.text == "云 OCR 提取文本"
    assert parsed.metadata["used_ocr"] is True
    assert parsed.metadata["ocr_provider"] == "cloud_openai_vision"


def test_pdf_parser_reports_ocr_disabled(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.base import DocumentOcrRequiredError
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrProviderError

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class DisabledProvider:
        def extract_pdf_text(self, file_bytes: bytes):
            raise OcrProviderError("该 PDF 需要 OCR 才能提取文本，但当前系统未启用 OCR")

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    with pytest.raises(DocumentOcrRequiredError, match="当前系统未启用 OCR"):
        PdfDocumentParser(ocr_provider=DisabledProvider()).extract(
            b"%PDF-1.4", "scan.pdf", "application/pdf"
        )


def test_pdf_parser_reports_empty_ocr_result(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.base import DocumentParseError
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrProviderError

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class EmptyProvider:
        def extract_pdf_text(self, file_bytes: bytes):
            raise OcrProviderError("OCR 已执行，但未提取到可用文本")

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    with pytest.raises(DocumentParseError, match="OCR 已执行，但未提取到可用文本"):
        PdfDocumentParser(ocr_provider=EmptyProvider()).extract(
            b"%PDF", "scan.pdf", "application/pdf"
        )


def test_pdf_parser_preserves_ocr_trace_fields(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.base import DocumentParseError
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrProviderError

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class BrokenProvider:
        def extract_pdf_text(self, file_bytes: bytes):
            raise OcrProviderError(
                "OCR 执行失败，请检查服务器 OCR 环境或稍后重试",
                stage="ocr_tesseract",
                error_code="ocr_tesseract_non_zero_exit",
                command="tesseract page.png stdout -l chi_sim+eng",
                exit_code=1,
                stderr="tesseract stderr detail",
            )

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())
    with pytest.raises(DocumentParseError) as exc_info:
        PdfDocumentParser(ocr_provider=BrokenProvider()).extract(
            b"%PDF", "scan.pdf", "application/pdf"
        )
    exc = exc_info.value
    assert getattr(exc, "stage", None) == "ocr_tesseract"
    assert getattr(exc, "error_code", None) == "ocr_tesseract_non_zero_exit"
    assert getattr(exc, "command", None) == "tesseract page.png stdout -l chi_sim+eng"
    assert getattr(exc, "exit_code", None) == 1
    assert getattr(exc, "stderr", None) == "tesseract stderr detail"
