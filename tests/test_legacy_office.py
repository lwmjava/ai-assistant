"""老 Office 格式（doc / xls / ppt）解析器测试。

覆盖：
- .xls 通过 xlrd 读取并按 sheet/行聚合文本
- .doc/.ppt 通过 LibreOffice 转换为现代格式后复用现有解析器
- LibreOffice 缺失 / 转换失败的错误语义
- 注册表与默认服务对老格式的调度
"""

from __future__ import annotations

import sys
import types
from pathlib import Path

import pytest


def test_xls_parser_extracts_sheet_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.rag.document_parsers.legacy_office import XlsDocumentParser

    class FakeSheet:
        name = "员工"
        nrows = 1

        def row_values(self, row_index: int) -> list[str]:
            return ["姓名", "Alice"]

    class FakeWorkbook:
        def sheets(self) -> list[FakeSheet]:
            return [FakeSheet()]

    fake_xlrd = types.ModuleType("xlrd")
    fake_xlrd.open_workbook = lambda file_contents=None: FakeWorkbook()
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

    parsed = XlsDocumentParser().extract(
        b"fake", "staff.xls", "application/vnd.ms-excel"
    )
    assert parsed.extension == "xls"
    assert "# Sheet: 员工" in parsed.text
    assert "Alice" in parsed.text
    assert parsed.metadata["parser_name"] == "xls"


def test_doc_parser_converts_and_reuses_docx_parser(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from docx import Document as Docx

    import app.rag.document_parsers.legacy_office as legacy

    path = tmp_path / "converted.docx"
    document = Docx()
    document.add_paragraph("老文档正文")
    document.save(path)
    docx_bytes = path.read_bytes()

    monkeypatch.setattr(
        legacy, "_convert_with_libreoffice", lambda file_bytes, source_ext, target_ext: docx_bytes
    )

    from app.rag.document_parsers.legacy_office import DocDocumentParser

    parsed = DocDocumentParser().extract(b"fake", "old.doc", "application/msword")
    assert parsed.extension == "doc"
    assert "老文档正文" in parsed.text
    assert parsed.metadata["parser_name"] == "doc"


def test_ppt_parser_converts_and_reuses_pptx_parser(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from pptx import Presentation

    import app.rag.document_parsers.legacy_office as legacy

    path = tmp_path / "converted.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "老幻灯片"
    presentation.save(path)
    pptx_bytes = path.read_bytes()

    monkeypatch.setattr(
        legacy, "_convert_with_libreoffice", lambda file_bytes, source_ext, target_ext: pptx_bytes
    )

    from app.rag.document_parsers.legacy_office import PptDocumentParser

    parsed = PptDocumentParser().extract(
        b"fake", "old.ppt", "application/vnd.ms-powerpoint"
    )
    assert parsed.extension == "ppt"
    assert "老幻灯片" in parsed.text
    assert parsed.metadata["parser_name"] == "ppt"


def test_ppt_parser_preserves_ocr_metadata_from_pptx_parser(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app.rag.document_parsers.legacy_office as legacy
    from app.rag.document_parsers.base import ParsedDocument

    monkeypatch.setattr(
        legacy, "_convert_with_libreoffice", lambda file_bytes, source_ext, target_ext: b"pptx-bytes"
    )

    def fake_extract(self, file_bytes: bytes, filename: str, content_type: str | None) -> ParsedDocument:
        return ParsedDocument(
            text="OCR 幻灯片文本",
            title="old",
            source=filename,
            extension="pptx",
            content_type=content_type,
            metadata={
                "parser_name": "pptx_ocr",
                "used_ocr": True,
                "ocr_provider": "tesseract",
            },
        )

    monkeypatch.setattr("app.rag.document_parsers.legacy_office.PptxDocumentParser.extract", fake_extract)

    from app.rag.document_parsers.legacy_office import PptDocumentParser

    parsed = PptDocumentParser().extract(
        b"fake", "old.ppt", "application/vnd.ms-powerpoint"
    )
    assert parsed.extension == "ppt"
    assert parsed.metadata["parser_name"] == "ppt_ocr"
    assert parsed.metadata["used_ocr"] is True
    assert parsed.metadata["ocr_provider"] == "tesseract"


def test_doc_parser_reports_missing_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    import app.rag.document_parsers.legacy_office as legacy
    from app.rag.document_parsers.base import DocumentParseError

    monkeypatch.setattr(legacy, "_libreoffice_binary", lambda: None)

    from app.rag.document_parsers.legacy_office import DocDocumentParser

    with pytest.raises(DocumentParseError, match="依赖缺失"):
        DocDocumentParser().extract(b"fake", "a.doc", "application/msword")


def test_doc_parser_reports_conversion_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    import app.rag.document_parsers.legacy_office as legacy
    from app.rag.document_parsers.base import DocumentParseError

    def boom(file_bytes: bytes, source_ext: str, target_ext: str) -> bytes:
        raise DocumentParseError("文件解析失败，请确认文件未损坏或未加密")

    monkeypatch.setattr(legacy, "_convert_with_libreoffice", boom)

    from app.rag.document_parsers.legacy_office import DocDocumentParser

    with pytest.raises(DocumentParseError, match="文件解析失败"):
        DocDocumentParser().extract(b"fake", "a.doc", "application/msword")


def test_registry_resolves_legacy_extensions() -> None:
    from app.rag.document_parsers.legacy_office import (
        DocDocumentParser,
        PptDocumentParser,
        XlsDocumentParser,
    )
    from app.rag.document_parsers.registry import DocumentParserRegistry

    registry = DocumentParserRegistry(
        [XlsDocumentParser(), DocDocumentParser(), PptDocumentParser()]
    )
    assert isinstance(registry.resolve("a.doc", None), DocDocumentParser)
    assert isinstance(registry.resolve("a.xls", None), XlsDocumentParser)
    assert isinstance(registry.resolve("a.ppt", None), PptDocumentParser)


def test_default_service_resolves_legacy_extensions() -> None:
    from app.rag.document_parsers.legacy_office import (
        DocDocumentParser,
        PptDocumentParser,
        XlsDocumentParser,
    )
    from app.rag.document_parsers.service import build_default_parser_service

    service = build_default_parser_service()
    assert isinstance(service._registry.resolve("a.doc", None), DocDocumentParser)
    assert isinstance(service._registry.resolve("a.xls", None), XlsDocumentParser)
    assert isinstance(service._registry.resolve("a.ppt", None), PptDocumentParser)
