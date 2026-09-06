"""RAG 异步导入平台测试。

覆盖：
- 批量文件导入任务创建与处理
- URL 导入与快照保存
- 失败任务重试
- 重解析生成版本链，并确保检索仅命中当前版本
"""

from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.api.deps import get_current_user
from app.core.config import settings
from app.core.security import Role
from app.main import app
from app.models.rag import Document, ImportBatch, ImportJob, ImportJobStatus
from app.models.user import User
from app.rag.import_jobs import run_import_jobs_once


def _patch_storage_root(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """将知识库源文件落盘目录重定向到临时路径。"""
    import app.rag.document_storage as storage_mod

    monkeypatch.setattr(storage_mod, "_PROJECT_ROOT", tmp_path)


@pytest.fixture()
def client():
    """关闭后台导入调度器，使用同步测试触发执行。"""
    original = settings.RAG_IMPORT_ENABLED
    settings.RAG_IMPORT_ENABLED = False
    tenant_id = f"import-tenant-{uuid.uuid4().hex}"
    user_id = f"import-user-{uuid.uuid4().hex}"
    fake_user = User(
        id=user_id,
        tenant_id=tenant_id,
        username="import-tester",
        hashed_password="",
        role=Role.TENANT_ADMIN.value,
        token_version=0,
        is_active=True,
    )
    app.dependency_overrides[get_current_user] = lambda: fake_user
    with TestClient(app) as c:
        yield c, fake_user
    app.dependency_overrides.clear()
    settings.RAG_IMPORT_ENABLED = original


@pytest.fixture()
def session():
    from app.core.database import engine, init_db

    init_db()
    with Session(engine) as s:
        yield s


def test_batch_upload_jobs_create_documents(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, fake_user = client
    _patch_storage_root(monkeypatch, tmp_path)

    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[
            ("files", ("one.txt", io.BytesIO("第一份知识".encode()), "text/plain")),
            ("files", ("two.txt", io.BytesIO("第二份知识".encode()), "text/plain")),
        ],
    )
    assert resp.status_code == 202
    payload = resp.json()
    assert payload["status"] in {"pending", "running"}
    assert len(payload["jobs"]) == 2

    batch_id = payload["id"]
    jobs = session.exec(
        select(ImportJob).where(ImportJob.batch_id == batch_id)
    ).all()
    assert len(jobs) == 2

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    batch = session.get(ImportBatch, batch_id)
    assert batch is not None
    assert batch.successful_jobs == 2
    assert batch.failed_jobs == 0

    docs = session.exec(
        select(Document).where(Document.tenant_id == fake_user.tenant_id, Document.is_current.is_(True))
    ).all()
    assert len(docs) >= 2


def test_url_import_and_reparse_create_version_chain(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    class FakeResponse:
        def __init__(self, text: str) -> None:
            self.content = text.encode("utf-8")
            self.headers = {"content-type": "text/html; charset=utf-8"}

        def raise_for_status(self) -> None:
            return None

    versions = [
        FakeResponse("<html><title>RAG</title><body>版本一 独有词 alpha</body></html>"),
        FakeResponse("<html><title>RAG</title><body>版本二 独有词 beta</body></html>"),
    ]

    async def fake_get(self, url: str, follow_redirects: bool = True):
        assert url == "https://example.com/rag"
        return versions.pop(0)

    monkeypatch.setattr("httpx.AsyncClient.get", fake_get)

    create = client.post(
        "/api/rag/import-jobs/url",
        json={"url": "https://example.com/rag"},
    )
    assert create.status_code == 202
    job_id = create.json()["id"]

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    first_job = session.get(ImportJob, job_id)
    assert first_job is not None
    assert first_job.status == ImportJobStatus.SUCCESS.value
    first_doc = session.get(Document, first_job.document_id)
    assert first_doc is not None
    assert first_doc.version_number == 1
    assert first_doc.is_current is True

    reparse = client.post(f"/api/rag/documents/{first_doc.id}/reparse")
    assert reparse.status_code == 202
    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    second_job = session.get(ImportJob, reparse.json()["id"])
    assert second_job is not None
    second_doc = session.get(Document, second_job.document_id)
    assert second_doc is not None
    assert second_doc.version_group_id == first_doc.version_group_id
    assert second_doc.version_number == 2
    assert second_doc.is_current is True

    session.refresh(first_doc)
    assert first_doc.is_current is False

    listing = client.get("/api/rag/documents")
    ids = [item["id"] for item in listing.json()]
    assert second_doc.id in ids
    assert first_doc.id not in ids


def test_retry_failed_url_job(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    async def failed_get(self, url: str, follow_redirects: bool = True):
        raise RuntimeError("network down")

    monkeypatch.setattr("httpx.AsyncClient.get", failed_get)
    create = client.post(
        "/api/rag/import-jobs/url",
        json={"url": "https://example.com/fail"},
    )
    assert create.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    failed_job = session.get(ImportJob, create.json()["id"])
    assert failed_job is not None
    assert failed_job.status == ImportJobStatus.FAILED.value

    class FakeResponse:
        content = b"<html><body>retry success</body></html>"
        headers = {"content-type": "text/html; charset=utf-8"}

        def raise_for_status(self) -> None:
            return None

    async def success_get(self, url: str, follow_redirects: bool = True):
        return FakeResponse()

    monkeypatch.setattr("httpx.AsyncClient.get", success_get)
    retried = client.post(f"/api/rag/import-jobs/{failed_job.id}/retry")
    assert retried.status_code == 202
    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    failed_job = session.get(ImportJob, failed_job.id)
    assert failed_job is not None
    assert failed_job.status == ImportJobStatus.SUCCESS.value
    assert failed_job.document_id is not None


def test_pdf_import_job_succeeds_with_ocr(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.rag.document_parsers.base import ParsedDocument

    def fake_parse(
        file_bytes: bytes, filename: str, content_type: str | None
    ) -> ParsedDocument:
        return ParsedDocument(
            text="OCR 后文本",
            title="scan",
            source=filename,
            extension="pdf",
            content_type=content_type,
            metadata={"parser_name": "pdf_ocr", "used_ocr": True},
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("scan.pdf", io.BytesIO(b"%PDF"), "application/pdf"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job_id = resp.json()["jobs"][0]["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.SUCCESS.value
    assert job.parser_name == "pdf_ocr"


def test_pdf_import_job_succeeds_with_cloud_ocr(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.rag.document_parsers.base import ParsedDocument

    def fake_parse(
        file_bytes: bytes, filename: str, content_type: str | None
    ) -> ParsedDocument:
        return ParsedDocument(
            text="云 OCR 后文本",
            title="scan",
            source=filename,
            extension="pdf",
            content_type=content_type,
            metadata={
                "parser_name": "pdf_ocr",
                "used_ocr": True,
                "ocr_provider": "cloud_openai_vision",
            },
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("scan.pdf", io.BytesIO(b"%PDF"), "application/pdf"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job_id = resp.json()["jobs"][0]["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.SUCCESS.value
    assert job.parser_name == "pdf_ocr"


def test_pdf_import_job_fails_when_ocr_environment_missing(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.rag.document_parsers.base import DocumentParseError

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        raise DocumentParseError("OCR 依赖缺失：未检测到 tesseract 或 chi_sim/eng 语言包")

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("scan.pdf", io.BytesIO(b"%PDF"), "application/pdf"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job = session.get(ImportJob, resp.json()["jobs"][0]["id"])
    assert job is not None
    assert job.status == ImportJobStatus.FAILED.value
    assert "OCR 依赖缺失" in (job.error or "")


def test_pdf_import_job_fails_when_cloud_ocr_http_error(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.rag.document_parsers.base import DocumentParseError

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        raise DocumentParseError("云 OCR 调用失败，请检查网络、鉴权或服务配置")

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("scan.pdf", io.BytesIO(b"%PDF"), "application/pdf"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job = session.get(ImportJob, resp.json()["jobs"][0]["id"])
    assert job is not None
    assert job.status == ImportJobStatus.FAILED.value
    assert "云 OCR 调用失败" in (job.error or "")


def test_pdf_import_job_persists_trace_for_ocr_provider_failure(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace
    from app.rag.document_parsers.pdf import PdfDocumentParser
    from app.rag.ocr.base import OcrProviderError

    class FakePage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    class BrokenProvider:
        def extract_pdf_text(self, file_bytes: bytes):
            raise OcrProviderError(
                "OCR 执行失败，请检查服务器 OCR 环境或稍后重试",
                stage="ocr_tesseract",
                error_code="ocr_tesseract_non_zero_exit",
                command="tesseract page.png stdout -l chi_sim+eng",
                exit_code=1,
                stderr="ocr stderr full detail",
            )

    monkeypatch.setattr("app.rag.document_parsers.pdf.PdfReader", lambda stream: FakeReader())

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        return PdfDocumentParser(ocr_provider=BrokenProvider()).extract(
            file_bytes, filename, content_type
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("scan.pdf", io.BytesIO(b"%PDF"), "application/pdf"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    job_id = resp.json()["jobs"][0]["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.FAILED.value
    assert "OCR 执行失败" in (job.error or "")
    assert "stderr:" in (job.error or "")

    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.stage == "ocr_tesseract"
    assert trace.error_code == "ocr_tesseract_non_zero_exit"
    assert trace.command == "tesseract page.png stdout -l chi_sim+eng"
    assert trace.exit_code == 1
    assert trace.stderr == "ocr stderr full detail"


def test_pptx_import_job_persists_trace_for_ocr_fallback_failure(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace
    from app.rag.document_parsers.office import PptxDocumentParser
    from app.rag.ocr.base import OcrProviderError

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = []
            self.has_notes_slide = False

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    class BrokenProvider:
        def extract_pdf_text(self, file_bytes: bytes):
            raise OcrProviderError(
                "OCR 执行失败，请检查服务器 OCR 环境或稍后重试",
                stage="ocr_tesseract",
                error_code="ocr_tesseract_non_zero_exit",
                command="tesseract slide.pdf stdout -l chi_sim+eng",
                exit_code=1,
                stderr="ppt ocr stderr detail",
            )

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)
    monkeypatch.setattr(
        "app.rag.document_parsers.office._convert_office_bytes_with_libreoffice",
        lambda file_bytes, source_ext, target_ext: b"%PDF-1.4 fake",
    )

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        return PptxDocumentParser(ocr_provider=BrokenProvider()).extract(
            file_bytes, filename, content_type
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[
            (
                "files",
                (
                    "deck.pptx",
                    io.BytesIO(b"fake-pptx"),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            )
        ],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    job_id = resp.json()["jobs"][0]["id"]
    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.stage == "pptx_pdf_ocr"
    assert trace.error_code == "ocr_tesseract_non_zero_exit"
    assert trace.command == "tesseract slide.pdf stdout -l chi_sim+eng"
    assert trace.exit_code == 1
    assert trace.stderr == "ppt ocr stderr detail"


def test_pptx_import_job_persists_trace_for_pdf_conversion_failure(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace
    from app.rag.document_parsers.libreoffice import LegacyOfficeConversionError
    from app.rag.document_parsers.office import PptxDocumentParser

    class FakeSlide:
        def __init__(self) -> None:
            self.shapes = []
            self.has_notes_slide = False

    class FakePresentation:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.slides = [FakeSlide()]

    monkeypatch.setattr("app.rag.document_parsers.office.Presentation", FakePresentation)

    def broken_convert(file_bytes: bytes, source_ext: str, target_ext: str) -> bytes:
        raise LegacyOfficeConversionError(
            message="PPT OCR 兜底失败：LibreOffice 转 PDF 失败",
            stage="legacy_office_convert",
            error_code="libreoffice_non_zero_exit",
            command="soffice --headless --convert-to pdf source.pptx",
            exit_code=1,
            stderr="libreoffice stderr detail",
        )

    monkeypatch.setattr(
        "app.rag.document_parsers.office._convert_office_bytes_with_libreoffice",
        broken_convert,
    )

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        return PptxDocumentParser().extract(file_bytes, filename, content_type)

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[
            (
                "files",
                (
                    "deck.pptx",
                    io.BytesIO(b"fake-pptx"),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            )
        ],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    job_id = resp.json()["jobs"][0]["id"]
    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.stage == "pptx_pdf_convert"
    assert trace.error_code == "libreoffice_non_zero_exit"
    assert trace.command == "soffice --headless --convert-to pdf source.pptx"
    assert trace.exit_code == 1
    assert trace.stderr == "libreoffice stderr detail"


def test_url_import_job_persists_trace_for_fetch_failure(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace

    async def failed_get(self, url: str, follow_redirects: bool = True):
        raise TimeoutError("should not be used directly")

    import httpx

    async def timeout_get(self, url: str, follow_redirects: bool = True):
        raise httpx.ReadTimeout("upstream timeout", request=httpx.Request("GET", url))

    monkeypatch.setattr("httpx.AsyncClient.get", timeout_get)
    create = client.post(
        "/api/rag/import-jobs/url",
        json={"url": "https://example.com/timeout"},
    )
    assert create.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job_id = create.json()["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.FAILED.value
    assert "URL 抓取失败：请求超时" in (job.error or "")

    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.stage == "url_fetch"
    assert trace.error_code == "url_timeout"
    assert trace.command == "GET https://example.com/timeout"
    assert "upstream timeout" in (trace.stderr or "")


def test_legacy_doc_import_job_succeeds(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from docx import Document as Docx

    path = tmp_path / "legacy.docx"
    document = Docx()
    document.add_paragraph("老文档导入正文")
    document.save(path)
    docx_bytes = path.read_bytes()

    import app.rag.document_parsers.legacy_office as legacy

    monkeypatch.setattr(
        legacy, "_convert_with_libreoffice", lambda file_bytes, source_ext, target_ext: docx_bytes
    )

    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("legacy.doc", io.BytesIO(b"%doc"), "application/msword"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()
    job_id = resp.json()["jobs"][0]["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.SUCCESS.value
    assert job.parser_name == "doc"


def test_legacy_doc_import_job_truncates_user_error_and_persists_trace(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace
    from app.rag.document_parsers.legacy_office import LegacyOfficeConversionError

    long_stderr = "LibreOffice detailed stderr. " * 60

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        raise LegacyOfficeConversionError(
            message="老 Office 格式转换失败：LibreOffice 执行超时",
            stage="legacy_office_convert",
            error_code="libreoffice_timeout",
            command="soffice --headless --convert-to docx source.doc",
            exit_code=None,
            stdout="",
            stderr=long_stderr,
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("legacy.doc", io.BytesIO(b"%doc"), "application/msword"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    job_id = resp.json()["jobs"][0]["id"]
    job = session.get(ImportJob, job_id)
    assert job is not None
    assert job.status == ImportJobStatus.FAILED.value
    assert job.error is not None
    assert "老 Office 格式转换失败" in job.error
    assert "stderr:" in job.error
    assert len(job.error) <= 500
    assert "LibreOffice detailed stderr." in job.error
    assert long_stderr not in job.error

    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.stage == "legacy_office_convert"
    assert trace.error_code == "libreoffice_timeout"
    assert trace.command == "soffice --headless --convert-to docx source.doc"
    assert trace.stderr == long_stderr
    assert trace.exception_type == "LegacyOfficeConversionError"


def test_legacy_doc_import_job_persists_full_trace_for_subprocess_failure(
    client: tuple[TestClient, User],
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    client, _ = client
    _patch_storage_root(monkeypatch, tmp_path)

    from app.models.rag import ImportJobTrace
    from app.rag.document_parsers.legacy_office import LegacyOfficeConversionError

    def fake_parse(file_bytes: bytes, filename: str, content_type: str | None):
        raise LegacyOfficeConversionError(
            message="老 Office 格式转换失败：LibreOffice 进程返回非 0 状态码",
            stage="legacy_office_convert",
            error_code="libreoffice_non_zero_exit",
            command="soffice --headless --convert-to pptx source.ppt",
            exit_code=1,
            stdout="stdout full context",
            stderr="stderr full context",
        )

    monkeypatch.setattr("app.rag.import_jobs.parse_uploaded_document", fake_parse)
    resp = client.post(
        "/api/rag/import-jobs/upload",
        files=[("files", ("legacy.ppt", io.BytesIO(b"%ppt"), "application/vnd.ms-powerpoint"))],
    )
    assert resp.status_code == 202

    import asyncio

    asyncio.run(run_import_jobs_once(limit=10))
    session.expire_all()

    job_id = resp.json()["jobs"][0]["id"]
    traces = session.exec(
        select(ImportJobTrace).where(ImportJobTrace.job_id == job_id)
    ).all()
    assert len(traces) == 1
    trace = traces[0]
    assert trace.exit_code == 1
    assert trace.stdout == "stdout full context"
    assert trace.stderr == "stderr full context"
    assert trace.message == "老 Office 格式转换失败：LibreOffice 进程返回非 0 状态码"


async def test_search_ignores_superseded_versions(
    session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    _patch_storage_root(monkeypatch, tmp_path)

    class FakeResponse:
        def __init__(self, text: str) -> None:
            self.content = text.encode("utf-8")
            self.headers = {"content-type": "text/html; charset=utf-8"}

        def raise_for_status(self) -> None:
            return None

    versions = [
        FakeResponse("<html><body>legacytoken 旧版本内容</body></html>"),
        FakeResponse("<html><body>newtoken 新版本内容</body></html>"),
    ]

    async def fake_get(self, url: str, follow_redirects: bool = True):
        return versions.pop(0)

    monkeypatch.setattr("httpx.AsyncClient.get", fake_get)

    from app.models.user import User
    from app.rag.import_jobs import create_reparse_job, create_url_import_job
    from app.rag.service import RAGService

    user = User(
        id=f"svc-user-{uuid.uuid4().hex}",
        tenant_id=f"svc-tenant-{uuid.uuid4().hex}",
        username="svc",
        hashed_password="",
        role=Role.TENANT_ADMIN.value,
        token_version=0,
        is_active=True,
    )

    job = create_url_import_job(session, user, url="https://example.com/versioned")
    await run_import_jobs_once(limit=10)
    session.expire_all()
    first = session.get(ImportJob, job.id)
    assert first is not None and first.document_id

    reparse = create_reparse_job(session, user, first.document_id)
    await run_import_jobs_once(limit=10)
    session.expire_all()
    second = session.get(ImportJob, reparse.id)
    assert second is not None and second.document_id

    rag = RAGService(session, user.tenant_id)
    results = await rag.search("legacytoken", top_k=5)
    assert all("legacytoken" not in row.content for row in results)
